"""
Builds the BA assessment Google Doc sources (docs/06) as Word files:
  assessment-kits/ba-docs/BA1-instructions.docx, BA1-answer-template.docx,
  BA2-instructions.docx, BA2-handoff-template.docx
Upload them to Google Drive and open with Google Docs (they convert), share "Anyone with the link
-> Viewer", then paste the links in /admin/stages. Each template's first line carries the marker
(CHASE-BA1 / CHASE-BA2) the platform checks for in a submitted copy: keep it.
Run: python3 scripts/ba-docs/build_docs.py
"""
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

OUT = Path(__file__).resolve().parents[2] / "assessment-kits" / "ba-docs"
GREY = RGBColor(0x6B, 0x72, 0x80)


def new_doc():
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(11)
    return d


def marker(d, text):
    p = d.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(8)
    r.font.color.rgb = GREY


def guide(d, text):
    p = d.add_paragraph()
    r = p.add_run(text)
    r.italic = True
    r.font.color.rgb = GREY
    return p


def para(d, text, bold_lead=None):
    p = d.add_paragraph()
    if bold_lead:
        p.add_run(bold_lead).bold = True
    p.add_run(text)
    return p


def bullets(d, items, style="List Bullet"):
    for it in items:
        d.add_paragraph(it, style=style)


def table(d, headers, rows=2, hint=None):
    t = d.add_table(rows=1 + rows, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        cell.paragraphs[0].add_run(h).bold = True
    if hint:
        for i, h in enumerate(hint):
            cell = t.rows[1].cells[i]
            cell.text = ""
            r = cell.paragraphs[0].add_run(h)
            r.italic = True
            r.font.color.rgb = GREY
    d.add_paragraph()
    return t


# ───────────────────────── BA 1: instructions ─────────────────────────
def ba1_instructions():
    d = new_doc()
    d.add_heading("BA Assessment 1: Solution and Spiky POV", 0)
    para(d, "About 2 hours of work. You have 3 hours from the moment you press Start on the assessment page.", "Time: ")

    d.add_heading("Why we're doing this", 1)
    para(d, "This is the job in miniature. Data from systems that don't talk to each other, a client who has asked for something that may not be what they need, and not much time. In the role you'd do exactly this in the first week of a client engagement, and then build the first version yourself.")
    para(d, "We want to see how you get from the mess to a clear, defensible solution, and how you explain it to the person who has to say yes. There is no single right answer. There are well-argued ones.")

    d.add_heading("The client and the ask", 1)
    para(d, "Kopano Connect is a mobile network dealer. Their Virtual Sales team phones existing business customers to renew and upgrade contracts.")
    para(d, "“Automate our renewal outreach. WhatsApp, SMS and email, starting three months before each contract ends, and give my agents one view per customer.”", "The ask, from Lerato Dube, GM Virtual Sales: ")

    d.add_heading("What you have", 1)
    bullets(d, [
        "kopano_vsam_extract.xlsx (download it from the assessment page after you press Start). Four sheets from four places that don't share a key: the monthly customer base export from the network's portal, one agent's own working spreadsheet, last month's telephony stats (the phone system only exports; it has no API), and the team's daily activity log.",
        "A 25-minute chat with Lerato on the “Interview the client” tab. She is busy and answers what you ask, not what you should have asked. Up to 25 messages. The chat is assessed.",
        "The open internet. AI tools are allowed and expected. Say what they did in your AI-use note.",
    ])

    d.add_heading("What to do", 1)
    bullets(d, [
        "On the assessment page, press “Make a copy of the template”. Google asks you to make a copy: do it. You must work in your own copy; only your copy is assessed.",
        "Go through the data. Find what's broken or missing before anyone tells you where to look.",
        "Interview Lerato. Find out what she hasn't told you.",
        "Fill in every section of your copy, replacing the grey guidance text as you go. Keep the template's first line.",
        "Share your copy: Share → General access → Anyone with the link → Viewer. Copy the link.",
        "Paste the link on the assessment page and press Submit before the timer runs out. We save a copy of your document at that moment; later edits aren't assessed.",
    ], "List Number")

    d.add_heading("What goes in your document", 1)
    bullets(d, [
        "Executive summary (at most 150 words): your recommendation and the decision you need from Lerato.",
        "The problem: the business question, what's in and out of scope.",
        "What the data shows: facts you can prove (sheet and column), and the insights that connect them.",
        "Spiky POV: one to three positions a reasonable person could disagree with, each with its evidence, the strongest counter-argument and why it loses here.",
        "Solution, architecture and tech stack: what to fix before building, what to build first, how data moves between the systems (given that some only export), what you'd use and why, and what you would not build.",
        "Success criteria: metric, baseline from the data, target, timeframe, and the result that would prove you wrong.",
        "Risks and open questions, and an AI-use note (at most 100 words).",
        "Appendix A: gap log. Appendix B: the questions you'd still ask.",
    ])
    para(d, "At most 1,500 words before Appendix A. Grey guidance text you leave in counts towards the limit.", "Limit: ")

    d.add_heading("What we're looking for", 1)
    bullets(d, [
        "What you found in the data, and how you prove it.",
        "What you got out of Lerato that she didn't volunteer.",
        "A point of view someone could disagree with, backed by evidence, with the best counter-argument answered.",
        "A solution and architecture that works with disconnected systems, messy data and a messy operation: practical, phased and honest about the constraints.",
        "Success criteria you could actually measure next month.",
        "Writing an executive can act on: answer first, numbers, a clear ask.",
    ])

    d.add_heading("The small print", 1)
    bullets(d, [
        "Your own work. Use AI however you like, but you must be able to explain everything in a live session.",
        "We will not use your work commercially, and you keep copyright.",
        "All the data is made up for this assessment. Any resemblance to a real company is a coincidence.",
        "If you need an adjustment (more time, a different format), ask through “Ask a person to review something” on your application page before you press Start.",
    ])
    d.save(OUT / "BA1-instructions.docx")


# ───────────────────────── BA 1: answer template ─────────────────────────
def ba1_template():
    d = new_doc()
    marker(d, "Chase Agents · BA Assessment 1 · answer template · CHASE-BA1 · keep this line")
    d.add_heading("Kopano Connect renewal outreach: solution and Spiky POV", 0)
    para(d, "", "Your name: ")
    guide(d, "How to use this template: replace every grey line with your answer. Keep the headings. Grey text you leave in counts towards the 1,500-word limit (everything before Appendix A). Tables can grow: add rows as you need them.")

    d.add_heading("1. Executive summary", 1)
    guide(d, "At most 150 words. Your recommendation in the first sentence, the two or three reasons that matter most (with numbers), and the decision you need from Lerato.")

    d.add_heading("2. The problem", 1)
    guide(d, "The business question in one sentence. Is the ask as stated the right problem? What is in scope and what is out, and why.")

    d.add_heading("3. What the data shows", 1)
    guide(d, "Facts first: specific and checkable, each cited to a sheet and column (or to the interview). Then the insights: patterns that only appear when you connect two or more facts or sources.")
    d.add_heading("Facts", 2)
    table(d, ["Fact", "Source (sheet / column, or interview)", "Why it matters"], 3,
          ["e.g. a count or share you calculated", "e.g. Base export, column …", "what it blocks or changes"])
    d.add_heading("Insights", 2)
    guide(d, "Two to four insights. Each one connects facts from more than one place and says what it means for the client.")

    d.add_heading("4. Spiky POV", 1)
    guide(d, "One to three positions a reasonable person could disagree with. Not “data quality matters”: something Lerato or a vendor would argue against.")
    for n in (1, 2):
        d.add_heading(f"POV {n}", 2)
        table(d, ["Position", "Evidence (facts and insights above)", "Strongest counter-argument", "Why it loses here"], 1)

    d.add_heading("5. Solution: architecture and tech stack", 1)
    d.add_heading("5.1 Fix before building", 2)
    guide(d, "What has to be true in the data or the operation before anything is automated, and who fixes it.")
    d.add_heading("5.2 What to build first, and what comes later", 2)
    guide(d, "Phase 1 in a few lines (what an agent and a manager can do on day one), then the later phases and what unlocks each one.")
    d.add_heading("5.3 Systems and data flow", 2)
    guide(d, "Which systems are involved and how data moves between them: what comes in by export or import, how often, and what becomes the single source of truth for a customer, a line and a contact point. A simple diagram is welcome (Insert → Drawing).")
    table(d, ["System", "What it holds", "How data gets in or out (API, export, manual)", "How often", "Source of truth for"], 4,
          ["e.g. Telephony platform", "", "", "", ""])
    d.add_heading("5.4 Tech stack and why", 2)
    table(d, ["Need", "What you'd use", "Why this, not the alternative", "Rough monthly cost (R)"], 4)
    d.add_heading("5.5 What not to build (yet), and why", 2)
    guide(d, "The things you are deliberately leaving out, including parts of what Lerato asked for, and what would change your mind.")

    d.add_heading("6. Success criteria", 1)
    guide(d, "Each one measurable from the data. The last column is the result that would prove your POV wrong.")
    table(d, ["Metric", "Baseline (from the data)", "Target", "By when", "Result that would prove you wrong"], 3)

    d.add_heading("7. Risks and open questions", 1)
    table(d, ["Risk", "Likelihood and impact", "What you'd do about it"], 3)

    d.add_heading("8. AI-use note", 1)
    guide(d, "At most 100 words: what AI tools did for you, how you checked their output, and what you did yourself.")

    d.add_heading("Appendix A: Gap log", 1)
    guide(d, "Every gap you found in the data, not just the ones in your summary.")
    table(d, ["Gap", "Evidence (sheet / column / row count)", "What it blocks", "Severity (Critical / High / Medium / Low)", "Proposed fix"], 6)

    d.add_heading("Appendix B: Questions you'd still ask", 1)
    table(d, ["Question", "Who you'd ask", "Why it matters"], 3)
    d.save(OUT / "BA1-answer-template.docx")


# ───────────────────────── BA 2: instructions ─────────────────────────
def ba2_instructions():
    d = new_doc()
    d.add_heading("BA Assessment 2: Build and handoff", 0)
    para(d, "About 4 hours of work in a 48-hour window that starts when you press Start on the assessment page.", "Time: ")
    d.add_heading("Why we're doing this", 1)
    para(d, "In the job you build the first working version and hand it to an engineer, who takes it to production. The handoff is the contract between you. We want to see a first version that proves the point, and a handoff an engineer could build from without asking you a single question.")
    d.add_heading("What you have", 1)
    bullets(d, [
        "The Solution Brief: the direction Lerato has agreed (on the assessment page).",
        "A cleaned dataset: customers, accounts, lines, contact points with consent flags, agents and normalised dates (download it after you press Start).",
        "Any free tools you like. Lovable or v0 with Supabase both work on free tiers.",
    ])
    d.add_heading("What to do", 1)
    bullets(d, [
        "Build a clickable MVP, hosted, running on the provided data: renewal queue, customer 360, log outcome (a callback date is required for “Call back”), template messaging blocked without consent, and a manager exceptions view.",
        "Press “Make a copy of the template” on the assessment page and write your handoff in your own copy. Keep its first line.",
        "Record a 5-minute Loom for Lerato: what it does, what it doesn't, and the decision you need from her.",
        "Share your handoff copy (Anyone with the link → Viewer). Paste the handoff link, the MVP link, the Loom link and its transcript on the assessment page, and submit. We save a copy of your handoff when you submit.",
    ], "List Number")
    d.add_heading("What we're looking for", 1)
    bullets(d, [
        "A data model with the right grain and keys (consent per contact point, history as events).",
        "An MVP that does what the brief asks, on the real data.",
        "A handoff with testable acceptance criteria, clear rules, an access matrix and edge cases. The test: could the engineer build without asking you?",
        "Judgement: what you left out, and why.",
        "A Loom an executive can act on.",
    ])
    d.add_heading("The small print", 1)
    bullets(d, [
        "Your own work; AI tools are allowed and expected.",
        "We will not use your work commercially, and you keep copyright.",
        "All the data is made up for this assessment.",
    ])
    d.save(OUT / "BA2-instructions.docx")


# ───────────────────────── BA 2: handoff template ─────────────────────────
def ba2_template():
    d = new_doc()
    marker(d, "Chase Agents · BA Assessment 2 · handoff template · CHASE-BA2 · keep this line")
    d.add_heading("Kopano Connect Renewal Desk: engineering handoff", 0)
    para(d, "", "Your name: ")
    para(d, "", "MVP link: ")
    guide(d, "Replace every grey line with your answer. Keep the headings. Write for the engineer who will build this without you in the room.")

    d.add_heading("1. The problem and the POV in five lines", 1)
    guide(d, "Five lines at most: the problem, your point of view, and what phase 1 must prove.")

    d.add_heading("2. Data model", 1)
    guide(d, "Every table, its grain in one sentence, and its keys. Add the ERD as a drawing here (Insert → Drawing) or as an extra file on the assessment page.")
    table(d, ["Table", "Purpose", "Grain (one row is…)", "Columns and types", "Keys and relationships"], 5)

    d.add_heading("3. User stories and acceptance criteria", 1)
    guide(d, "One row per story. Acceptance criteria in Given / When / Then, specific enough to test.")
    table(d, ["ID", "As a … I want … so that …", "Acceptance criteria (Given / When / Then)"], 5)

    d.add_heading("4. Business rules", 1)
    guide(d, "Eligibility for the renewal queue, consent and opt-out, allocation to agents, deduplication. Each rule precise enough to code.")
    table(d, ["Rule", "Exactly how it works", "Example"], 4)

    d.add_heading("5. Access matrix", 1)
    guide(d, "Who can see and do what. The engineer enforces this in the database, so be exact.")
    table(d, ["Action", "Agent", "Manager", "Admin"], 6)

    d.add_heading("6. Edge cases", 1)
    table(d, ["Situation", "What should happen"], 4)

    d.add_heading("7. Non-functional needs", 1)
    guide(d, "Volumes, speed, audit, privacy (POPIA), availability: whatever the engineer must design for.")

    d.add_heading("8. Out of scope, and why", 1)
    guide(d, "What phase 1 deliberately doesn't do, and what would bring it in.")

    d.add_heading("9. Open questions", 1)
    table(d, ["Question", "Who answers it", "What it blocks"], 3)
    d.save(OUT / "BA2-handoff-template.docx")


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    ba1_instructions()
    ba1_template()
    ba2_instructions()
    ba2_template()
    print("written to", OUT)
